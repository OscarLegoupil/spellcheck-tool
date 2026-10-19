"""Build matching English and Finnish technical operations DOCX examples."""

import re
from collections import Counter
from pathlib import Path
from runpy import run_path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ENGLISH_SOURCE = run_path(str(ROOT / "scripts" / "build_english_technical_example.py"))
PAGES = ENGLISH_SOURCE["PAGES"]
SUPPLEMENTAL_PARAGRAPHS = ENGLISH_SOURCE["SUPPLEMENTAL_PARAGRAPHS"]
TYPO_LEDGER = ENGLISH_SOURCE["TYPO_LEDGER"]


ENGLISH_OUTPUT = ROOT / "examples" / "northstar-edge-operations-en.docx"
FINNISH_OUTPUT = ROOT / "examples" / "northstar-edge-operations-fi.docx"
NAVY = "19334F"
TEAL = "007A7A"
PALE = "EAF3F3"
STRIPE = "F2F5F7"
FINNISH_TYPO_LEDGER = [
    "arkkitehturi",
    "kuukausitainen",
    "hyvaksytty",
    "aluelinen",
    "ensijainen",
    "varajarjestelma",
    "pyynto",
    "turvalisuus",
    "todenus",
    "kaytoikeus",
    "saluaus",
    "tuniste",
    "valovnta",
    "mitaus",
    "kynnyasarvo",
    "halytys",
    "vieive",
    "vikasietosuus",
    "toissijanen",
    "julkiaisu",
    "kayttoonotto",
    "epaonnistuminen",
    "palautusmenetely",
    "tietoliikennee",
    "kuormtus",
    "virheprosenti",
    "mittaristoo",
    "varmuuskopo",
    "sailytys",
    "kayttaja",
    "paakayttaja",
    "kaynnistys",
    "paivitys",
    "hairio",
    "yhdeys",
    "yllapito",
    "sahkoposti",
    "kayttotunnus",
    "lokitus",
    "kayttajaohje",
]

FINNISH_PAGES = [
    {
        "title": "Reunayhdyskäytävä",
        "subtitle": "Käyttö- ja luotettavuuskäsikirja | Versio 3.2 | 16.10.2026",
        "items": [
            ("p", "Reunayhdyskäytävä välittää alueelliset tuotepalvelujen pyynnöt taustapalveluille. Se päättää asiakkaan suojatun yhteyden, soveltaa reitityssääntöjä ja välittää pyynnöt itsenäisesti julkaistuille palveluille. Sen arkkitehturi on tarkoituksella pieni: kolme tilatonta yhdyskäytäväryhmää, ohjaustaso ja hallittu avainsäilö. Tämä käsikirja kuvaa tuotantokäytäntöjä sovelluskehityksen sijaan."),
            ("callout", "PALVELUTAVOITE", "Palvelun kuukausitainen saatavuustavoite on 99,95 %. Pyyntö lasketaan saatavaksi, kun yhdyskäytävä palauttaa kelvollisen vastauksen tai dokumentoidun taustavirheen kahdessa sekunnissa. Tavoitteen ulkopuolelle jää huolto, josta on ilmoitettu vähintään seitsemän päivää etukäteen."),
            ("p", "Ensisijainen tuotantoalue palvelee asiakkaita, ja lämmin vara-alue pidetään valmiina hallittua palautusta varten. Aluetiimit vastaavat kapasiteetista, varmenteista ja häiriötilanteista. Käyttölista on tarkoitettu kahden henkilön muutosikkunaan: toinen toteuttaa ja toinen tarkastaa. Tuotantoon ei tehdä muutoksia jäljittämättömästä paikallisesta tiedostosta."),
            ("table", ["Palvelun ominaisuus", "Tuotannon perustaso", "Vastuu"], [["Saatavuustavoite", "99,95 % kalenterikuukaudessa", "SRE"], ["Tavallinen pyyntöraja", "2 sekuntia päästä päähän", "Alusta"], ["Suurin pyyntökoko", "8 megatavua; suuremmat hylätään", "Rajapintatiimi"], ["Ulkoinen riippuvuus", "Avainsäilö ja alueellinen nimipalvelu", "Tietoturva"]]),
            ("p", "Yhdyskäytävä ei säilytä liiketoiminnan tietueita. Se voi pitää lyhytaikaisen pyynnön muistissa välityksen ajan; käyttöloki sisältää tunnisteet ja ajoitukset, ei pyyntösisältöä. Tiimit tarkastavat aluelinen kapasiteetin ennen liikenteen siirtoa."),
        ],
        "extra": ["Kapasiteettia arvioidaan kiireisimmän viiden minuutin jakson, samanaikaisten yhteyksien ja muistinkäytön perusteella. Kuukausikatsaus sisältää kahden viikon kasvuennusteen, jotta alueelliset kiintiöt ehditään varata ennen kuormitushuippua.", "Palveluomistaja julkaisee vain hyvaksytty reitityspaketin. Jokaiselle uudelle reitille dokumentoidaan tavoitenopeus, tietoluokka ja taustapalvelun vastuutiimi."]
    },
    {
        "title": "Arkkitehtuuri ja pyyntövirta",
        "subtitle": "Liikennepolku, vikavyöhykkeet ja kapasiteetin rajat",
        "items": [
            ("p", "Julkinen nimipalvelu ohjaa alueellisen nimen hallittuun kuormantasaajaan. Se jakaa liikenteen kolmeen yhdyskäytäväryhmään eri vika-alueilla. Yhdyskäytävät noutavat allekirjoitetun reitityspaketin ohjaustasolta ja välittävät pyynnöt todennetuilla yhteyksillä. Rakenteen ensijainen hyöty on se, että kapasiteettia voidaan kasvattaa muuttamatta reitityspolitiikkaa."),
            ("diagram", [("Asiakas", "Salattu yhteys\nNimipalvelu"), ("Alueellinen LB", "Terveystarkistukset\nVyöhyketietoinen"), ("Yhdyskäytävät", "3 vyöhykettä\nTilattomat"), ("Palvelut", "Yksityisreitit\nReittikohtaiset rajat")]),
            ("table", ["Vaihe", "Tyypillinen viive", "Vikamerkki"], [["Nimihaku ja yhteys", "20–80 millisekuntia", "Nimihaku tai kättely epäonnistuu"], ["Kuormantasaaja", "5–15 millisekuntia", "Alle 2 vyöhykettä terveenä"], ["Reitityspolitiikka", "Alle 10 millisekuntia", "Paketin ikä yli 15 minuuttia"], ["Taustapalvelu", "Reittikohtainen", "Aikakatkaisu tai 5xx-kasvu"]]),
            ("p", "Yhdyskäytävä mittaa asiakkaan havaitseman vasteajan erikseen taustapalvelun kestosta. Kapasiteettihälytys käyttää kokonaisviivettä, koska se sisältää jonotuksen ja yhteyden muodostamisen. Normaaliin kuormaan sisältyy yksi ylimääräinen instanssi ryhmää kohden."),
            ("callout", "VIKAVYÖHYKE", "Toissijanen alue on itsenäinen: sillä on omat nimipalvelutietueet, varmenteet ja palvelutunnukset. Vara-alue ei jaa ensisijaisen alueen verkkoreittiä eikä julkaisulukkoa."),
            ("p", "Terveystarkistus mittaa prosessin valmiuden ja aktiivisen reitityspaketin. Se ei kutsu kaikkia taustapalveluja, jotta yhden hitaan palvelun häiriö ei poistaisi muuten toimivaa kapasiteettia."),
        ],
        "extra": ["Reitit arvioidaan järjestyksessä: nimi, polku, menetelmä ja lopuksi taustapalvelu. Jokainen pyynto saa vakaan virhekoodin, eikä turvalisuus vaarannu liian yleisen reitin vuoksi. Vara-alueen varajarjestelma pidetään valmiina.", "Reitityspaketissa on kasvava versionumero, luontiaika ja hyväksyvän omistajan tunniste. Yhdyskäytävä hylkää virheellisen allekirjoituksen tai tuntemattoman rakenteen ja jatkaa viimeisellä varmennetulla paketilla."]
    },
    {
        "title": "Asennus ja käynnistys",
        "subtitle": "Solmukuvan valmistelu ja liittäminen alueelliseen ryhmään",
        "items": [
            ("p", "Solmukuvat rakennetaan hyväksytystä peruskuvasta julkaisuketjussa. Käynnistys tarkistaa kuvan allekirjoituksen, asentaa yhdyskäytäväpaketin ja pyytää lyhytikäisen solmutunnisteen. Kuva on muuttumaton julkaisun jälkeen; asetukset liitetään käynnistyksessä eikä niitä korjata käsin."),
            ("table", ["Vaihe", "Vaadittu näyttö", "Keskeytysehto"], [["Kuvan rakennus", "Allekirjoitettu tiiviste", "Allekirjoitus virheellinen"], ["Solmun rekisteröinti", "Vyöhyke ja omistajamerkinnät", "Tunnistepyyntö hylätään"], ["Valmius", "Paketti ladattu, tarkistus läpäisty", "Vanhentunut politiikka"], ["Liikenteen avaus", "Mittarit perustasolla", "Virheaste tai viive kasvaa"]]),
            ("bullets", ["Varmista hyväksytty aliverkko ja oikea käytettävyysalue.", "Tarkista ohjaustason yhteys ennen kuormantasaajan kohteen avaamista.", "Pidä kuvan julkaisu ja liikenteen avaus erillisinä hyväksyttyinä vaiheina."]),
            ("p", "Kuva ei sisällä ympäristökohtaisia salaisuuksia. Käynnistyksessä työkuormatunniste vaihtaa solmun väitteen lyhytikäiseen varmenteeseen ja noutaa allekirjoitetun reitityspaketin. Sama artefakti voidaan siirtää testi-, vaiheistus- ja tuotantoympäristöihin uudelleenrakentamatta."),
            ("callout", "PALAUTUS", "Jos käynnistys epäonnistuu rekisteröinnin jälkeen, poista solmu liikenteestä, säilytä käynnistysloki ja korvaa se viimeisellä toimivalla kuvalla. Käsin tehty korjaus tuottaa tilan, jota ei voi toistaa ylläpidon aikana."),
        ],
        "extra": ["Tarkista käyttöjärjestelmän päivitystaso, valvonta-agentin tila ja kuvan tiiviste ennen rekisteröintiä. Käynnistysloki arkistoidaan solmun tunnisteella, jotta pakettivirhe voidaan erottaa verkkohäiriöstä.", "Uusi solmu otetaan liikenteeseen vaiheittain: ensin synteettinen pyyntö, sitten pieni asiakasosuus. Jos mittari poikkeaa normaalista, pysäytä avaaminen ja liitä loki muutoslippuun."]
    },
    {
        "title": "Verkko ja liikennepolitiikka",
        "subtitle": "Sisääntulo, ulosmenevä liikenne, aikarajat ja muutokset",
        "items": [
            ("p", "Sisääntulo sallitaan vain alueelliselta kuormantasaajalta julkaistuun salattuun porttiin. Yhdyskäytävässä ei ole julkista hallintaporttia. Ulosmenevä liikenne estetään oletusarvoisesti ja avataan vain nimetyille taustapalveluille, ohjaustasolle, nimipalvelulle ja kellon synkronoinnille. Verkkomuutokset vaativat vertaisarvion, koska laaja sääntö voi ohittaa palvelurajat."),
            ("table", ["Politiikka", "Oletus", "Muutosohje"], [["Asiakkaan aikaraja", "2 sekuntia", "Nosta vain mitattuun tarpeeseen"], ["Taustapalvelun aikaraja", "1,5 sekuntia", "Alita kutsujan määräaika"], ["Uusintayritykset", "1 turvalliselle lukupyynnölle", "Älä toista kirjoittavaa pyyntöä"], ["Yhteyspooli", "Reitti- ja aluekohtainen", "Tyhjennä ennen pienentämistä"], ["Ulosmenevä liikenne", "Estetty oletuksena", "Palvelu, portti ja päättymisaika"]]),
            ("p", "Reititysmuutos julkaistaan allekirjoitettuna versiona. Ohjaustaso tarkistaa rakenteen, omistajuuden ja aikarajojen rajat ennen jakelua. Julkaistua versiota ei muokata käsin, vaan virhe korjataan lähdekoodissa ja julkaistaan uutena versiona."),
            ("bullets", ["Käytä yksiselitteistä ympäristömerkintää.", "Testaa uusi taustareitti koesolmussa.", "Varmista alueellisen nimipalvelun kohde muutosarviossa."]),
            ("p", "Verkkosääntö sisältää teknisen omistajan, käyttötarkoituksen ja tarkistuspäivän. Vanhentunut poikkeus poistetaan automaattisesti, ellei omistaja uusi sitä hyväksyntäpolun kautta. Kiireellinen avaus rajataan ajallisesti ja tarkastetaan seuraavana työpäivänä."),
        ],
        "extra": ["Reittikohtainen nopeusraja asetetaan asiakkaan tunnisteen perusteella, ja alueellinen kokonaisraja toimii viimeisenä suojana. Omistaja ilmoittaa odotetun kuormituspiikin ennen rajan nostamista.", "Aikarajat muodostavat yhteisen budjetin. Asiakkaan määräajan on jätettävä aikaa yhdyskäytävän ja taustapalvelun käsittelylle sekä hyödyllisen virheen palauttamiselle."]
    },
    {
        "title": "Tunnistus ja salaisuuksien käsittely",
        "subtitle": "Työkuormatunnisteet, ylläpitäjät ja varmenteiden kierto",
        "items": [
            ("p", "Henkilöstön käyttö perustuu yrityksen tunnistuspalveluun ja vahvaan monivaiheiseen tunnistukseen. Ylläpitäjä saa määräaikaisen käyttöoikeuden alueelliseen ohjaustasoon nimetyn ryhmän kautta. Jaetut tunnukset ovat kiellettyjä. Toinen tarkastaja hyväksyy tuotantomuutoksen."),
            ("table", ["Tunniste", "Voimassaolo", "Säilytys ja kierto"], [["Solmutunniste", "15 minuuttia", "Uusitaan työkuormatunnisteella"], ["Yhdyskäytävän varmenne", "30 päivää", "Automaattinen kierto 10 päivää ennen"], ["Ylläpitäjän istunto", "8 tuntia", "Uusi tunnistus oikeustoimille"], ["Hätätunniste", "60 minuuttia", "Holvattu, auditoitu ja peruttu käytön jälkeen"]]),
            ("p", "Palvelu saa vain tehtävän vaatiman käyttöoikeuden. Solmu voi lukea reitityspaketin ja ilmoittaa terveystiedon, mutta se ei voi muuttaa politiikkaa eikä noutaa asiakastietoja. Henkilötunnisteita ei kopioida kuviin, käynnistystiedostoihin tai komentohistoriaan."),
            ("callout", "VARMENTEEN KIERTO", "Julkaise uusi luottamusketju ensin, varmista sen toimivuus kaikissa alueissa ja poista vanha avain vasta onnistuneiden kättelymittareiden jälkeen."),
            ("p", "Salaisuuksien käyttö kirjataan toimijan, resurssin, alueen ja pyynnön tunnisteella. Lokit eivät sisällä salaisia arvoja. Jos tunniste voi olla paljastunut, peru se heti ja säilytä tarkastusketju. Liikenne salataan siirron aikana ja avainsäilö suojaa avaimet."),
        ],
        "extra": ["Ylläpitäjän kaytoikeus myönnetään vasta tunnistuksen todenus jälkeen. Jokaisella työkuorman tuniste on oltava omistaja.", "Varmenteen saluaus ja kiertoa valvotaan sekä päättymispäivän että yhteyden onnistumisen perusteella. Uusinnan epäonnistuminen ilmoitetaan tietoturvan päivystäjälle ennen päällekkäisyysjakson päättymistä."]
    },
    {
        "title": "Havainnointi ja palvelutavoitteet",
        "subtitle": "Mittarit, kynnykset ja viikoittainen luotettavuuskatsaus",
        "items": [
            ("p", "Jokainen pyyntö tuottaa keston, vastausluokan, reitin tunnisteen ja alueen. Pyyntösisältöä tai valtuutusotsakkeita ei tallenneta. Havainnointijärjestelmä kokoaa tiedot alueittain ja palveluittain; mittaristo näyttää asiakkaan vasteajan ja yhdyskäytävän jonon pituuden."),
            ("table", ["Signaali", "Varoitus", "Välitön hälytys"], [["Saatavuus", "Alle 99,97 % / 30 minuuttia", "Alle 99,90 % / 10 minuuttia"], ["p95 vasteaika", "Yli 700 millisekuntia / 15 minuuttia", "Yli 1,2 sekuntia / 5 minuuttia"], ["Yhdyskäytävän 5xx", "Yli 0,5 % / 10 minuuttia", "Yli 2 % / 5 minuuttia"], ["Paketin ikä", "Yli 5 minuuttia", "Yli 15 minuuttia"], ["Terveet alueet", "Kaksi jäljellä", "Yksi jäljellä"]]),
            ("p", "Hälytyskynnys ei ole lupa odottaa, jos asiakashaitta on jo vahvistettu. Päivystäjä kirjaa ensimmäisen tunnetun vaikutusajan, tarkistaa alueellisen tilanteen ja tunnistaa eniten kärsivän reitin. Muutosikkunan aikana mittausdataa verrataan edellisen viikon vastaavaan aikaan."),
            ("bullets", ["Tarkista alueen tila, reitti ja vyöhyke.", "Vertaa yhdyskäytävän ja taustapalvelun kestoa.", "Merkitse julkaisut ja verkkohuollot mittaristoon."]),
            ("callout", "TIEDON LAATU", "Jos mittaus näyttää väärän yksikön tai aukon, korjaa mittausketju äläkä nosta kynnysarvoa ongelman peittämiseksi. Jokainen hälytys kirjataan ja sen omistaja nimetään."),
        ],
        "extra": ["Valovnta seuraa vasteajan mitaus ja kynnyasarvo on kirjattu mittaristoon. Jokainen halytys saa nimetyn omistajan.", "Palvelutason mittari lasketaan onnistuneista pyynnöistä suhteessa kaikkiin hyväksyttyihin pyyntöihin. Rajaukset dokumentoidaan mittarin yhteydessä. Viikkokatsaus arvioi, onko hälytys asiakkaalle merkityksellinen ja voiko omistajatiimi toimia sen perusteella. Reittinimet tulevat hyväksytystä rekisteristä ja tuntemattomat arvot näkyvät erillisessä luokassa."]
    },
    {
        "title": "Tietojen kestävyys ja aluepalautus",
        "subtitle": "Tilattomuus, varmistukset ja liikenteen palauttaminen",
        "items": [
            ("p", "Yhdyskäytävä on tilaton. Reitityspaketit ja varmenteet säilyvät hallituissa järjestelmissä, kun taas solmun paikallinen välimuisti voidaan hylätä. Solmu lataa viimeisimmän hyväksytyn paketin käynnistyessään ja ilmoittaa version ennen liikenteen avaamista."),
            ("table", ["Aineisto", "Palautuslähde", "Tavoite"], [["Reitityspaketti", "Allekirjoitettu versiohistoria", "RPO 0; palautus 10 min"], ["Varmenteet", "Hallittu avainsäilö", "RPO 0; kierto 30 min"], ["Solmun tila", "Muuttumaton kuvarekisteri", "Uudelleenrakennus 20 min"], ["Tarkastustapahtumat", "Alueellinen lokiarkisto", "RPO alle 5 min"]]),
            ("p", "Ensisijainen ja vara-alue käyttävät samaa hyväksyttyä politiikkaversiota, mutta niiden verkko ja laskenta ovat erillisiä. Yhdenmukaisuus varmistetaan pakettien tiivisteillä; onnistunut julkaisu ei yksin todista, että jokainen vyöhyke sai uuden version."),
            ("bullets", ["Tarkista vara-alueen palvelut ja varmenteet.", "Vertaa pakettien tiivisteitä ja taustapalvelujen kapasiteettia.", "Siirrä liikennettä vaiheittain ja pysäytä virheiden kasvaessa."]),
            ("p", "Solmu voidaan luoda uudelleen kuvarekisteristä ilman tietokantapalautusta. Kadonnut politiikkaversio haetaan allekirjoitetusta historiasta. Jos hallittu riippuvuus ei ole käytettävissä, pidä liikenne terveellä alueella ja sovi palautuksesta häiriönjohtajan kanssa."),
            ("callout", "PALAUTUKSEN TARKISTUS", "Testaa edustava luku- ja kirjoitusreitti, varmista tarkastustapahtumien saapuminen ja pidä aiempi alue valmiina, kunnes mittarit ovat vakaat 30 minuutin ajan."),
        ],
        "extra": ["Palvelun vieive tarkistetaan ennen siirtoa. Vara-alueen vikasietosuus testataan kuukausittain, ja vara-alue pidetään valmiina julkiaisu varten.", "Aluepalautuksesta tehdään päätöstietue, jossa nimetään häiriönjohtaja, kohdealue ja siirron syy. Tietue sisältää jokaisen liikenneosuuden sekä tarkistusajan. Paluu alkuperäiselle alueelle on erillinen muutos. Omistaja varmistaa vian korjauksen, ajantasaisen politiikkapaketin ja vakaat mittarit ennen paluuta."]
    },
    {
        "title": "Julkaisut ja muutostenhallinta",
        "subtitle": "Koesiirrot, palautusrajat ja tuotannon näyttö",
        "items": [
            ("p", "Julkaisu siirtyy kehityksestä vaiheistukseen ja tuotantoon samasta allekirjoitetusta artefaktista. Muutostietue sisältää lähdeversion, kuvan tiivisteen, riippuvuusraportin, politiikkarakenteen ja hyväksyjät. Julkaisu päättyy vasta koejakson ja mittarinäytön tarkastuksen jälkeen."),
            ("table", ["Vaihe", "Liikenneosuus", "Vähimmäisseuranta"], [["Koe", "1 % yhdessä vyöhykkeessä", "10 minuuttia"], ["Alueellinen", "10 % vyöhykkeissä", "20 minuuttia"], ["Laaja", "50 %, sitten 100 %", "30 min / vaihe"], ["Päätös", "Kaikki terveet kohteet", "Jälkitarkastus"]]),
            ("p", "Julkaisun omistaja seuraa mittaristoa ja tarkastaja muutoslippua. Keskeytä, jos virhebudjetin kulutus kaksinkertaistuu, pakettitiiviste poikkeaa tai vyöhyke menettää valmiutensa. Palautus käyttää viimeistä toimivaa kuvaa ja politiikkaversiota."),
            ("bullets", ["Varmista asiakkaan ja taustapalvelun tila ennen liikenteen siirtoa.", "Kirjaa kuvan ja politiikan versio jokaisessa vaiheessa.", "Testaa API-yhteensopivuus synteettisellä pyynnöllä."]),
            ("callout", "ONNISTUMISRAJA", "Julkaisu suljetaan vasta, kun kaikki alueet ilmoittavat oikean version, palautusikkuna on dokumentoitu eikä uusi vakava hälytys ole ilman omistajaa."),
        ],
        "extra": ["Kayttoonotto alkaa vasta, kun artefakti on tarkistettu. Epaonnistuminen käynnistää palautusmenetely ja tietoliikennee tarkistuksen.", "Rakennusketju tallentaa allekirjoitetun kuvan rinnalle lähdeversion, rakentajan tunnisteen ja riippuvuusluettelon. Tuotantosolmu hyväksyy vain allekirjoitetun artefaktin hyväksytystä rekisteristä. Jos koejulkaisu pysäytetään, kirjaa mittari ja havainto, joka esti jatkamisen. Korjattu ehdokas saa uuden version ja käy vaaditut vaiheet uudelleen."]
    },
    {
        "title": "Häiriötilanteiden hallinta",
        "subtitle": "Rajaus, eristäminen, viestintä ja oppiminen",
        "items": [
            ("p", "Ensimmäinen vastaaja kuittaa hälytyksen, avaa häiriökanavan ja nimeää häiriönjohtajan. Johtaja koordinoi päätökset, tekninen vastuuhenkilö tutkii vikaa ja kirjuri ylläpitää aikajanaa. Pienessä tilanteessa sama henkilö voi hoitaa useita tehtäviä, mutta vastuut sovitaan ennen riskialtista muutosta."),
            ("table", ["Vakavuus", "Asiakasvaikutus", "Ensivaste"], [["SEV-1", "Laaja katko tai tietovuoto", "Välitön päivystys ja tilannehuone"], ["SEV-2", "Alue- tai reittihäiriö", "Päivystys, arvio 10 minuutissa"], ["SEV-3", "Rajattu haitta, kiertotapa", "Omistajan vastaus työaikana"], ["SEV-4", "Ei aktiivista asiakashaittaa", "Normaali työjono"]]),
            ("p", "Liikennevian selvitys alkaa alueen saatavuudesta, yhdyskäytävän 5xx-määrästä ja taustapalvelun viiveestä. Tarkista viimeisimmät julkaisut ennen solmujen uudelleenkäynnistystä. Tietoturvaepäilyssä peru tunniste ensin ja säilytä käyttöloki. Asiakastietoja tai salaisuuksia ei lisätä häiriökanavaan."),
            ("bullets", ["Kerro mikä tiedetään, mikä on epävarmaa ja mitä testataan.", "Nimeä asiakasviestinnän ja teknisen aikajanan omistajat.", "Nosta asia palveluomistajalle, jos vaikutus ylittää alueen."]),
            ("p", "Häiriönjohtaja vastaa riskipäätöksistä ja palautuksen tarkastuspisteestä. Vakauttamisen jälkeen tallennetaan vaikutusikkuna, syy-yhteydet ja jatkotoimien omistajat. Katsaus keskittyy järjestelmän tapahtumiin, ei syyllisten etsimiseen."),
            ("callout", "VIESTINTÄ", "SEV-1-tilanteessa sisäinen päivitys julkaistaan 15 minuutin välein. Asiakasviestissä kerrotaan myös seuraavan päivityksen ajankohta, vaikka uutta diagnoosia ei vielä olisi."),
        ],
        "extra": ["Kuormtus ja virheprosenti näkyvät valvonnassa. Mittaristoo tarkistetaan ennen palautuspäätöstä, ja varmuuskopo säilytetään erillisessä arkistossa.", "Paivitys katkaisee yhdeys vain hyväksytyllä hetkellä. Hairio ja yllapito kirjataan aikajanaan. Häiriötilanne dokumentoidaan palvelun palautuksen jälkeen."]
    },
    {
        "title": "Huolto ja viiteohjeet",
        "subtitle": "Toistuvat tehtävät, säilytysajat ja muutosikkuna",
        "items": [
            ("p", "Huolto sovitaan aluekalenterissa ja yhdistetään muutostietueeseen. Omistaja vahvistaa asiakasvaikutuksen, versiot, palautusohjeen ja henkilöstön ennen ikkunan alkua. Huolto ei saa osua suuren julkaisun, riippuvuushäiriön tai aluepalautusharjoituksen kanssa päällekkäin."),
            ("table", ["Tiheys", "Tehtävä", "Näyttö"], [["Päivittäin", "Tarkista hälytykset ja terveystarkistukset", "Päivystäjän vuoromuistio"], ["Viikoittain", "Tarkista kapasiteetti ja politiikkaerot", "Luotettavuuskatsaus"], ["Kuukausittain", "Testaa varmenteen kierto ja palautus", "Harjoitustietue"], ["Neljännesvuosittain", "Tarkasta käyttöoikeudet ja ulosmenevät säännöt", "Omistajan vahvistus"], ["Vuosittain", "Harjoittele aluepalautus", "Toimintasuunnitelma"]]),
            ("p", "Tarkastustapahtumia säilytetään 400 päivää. Mittarit pidetään täydellä tarkkuudella 30 päivää ja harvennettuina vuoden ajan. Säilytys tukee vaatimustenmukaisuutta ja rajoittaa tarpeetonta tiedon keruuta. Arkistovienti salataan ja sen palautus testataan vuosiharjoituksessa."),
            ("bullets", ["Varmista muutoslippu, vertaisarvio ja päivystysvuoro.", "Tallenna alkutila, kuvan tiiviste, politiikkaversio ja liikenneosuus.", "Suorita jälkitarkistukset ja liitä näyttö muutokseen."]),
            ("callout", "TARKASTUSTIETUE", "Tietue sisältää toimijan, ajan, alueen, vanhan ja uuden version, hyväksynnän sekä tarkistustuloksen. Säilytä se muutosjärjestelmässä, älä henkilökohtaisessa muistiinpanossa."),
        ],
        "extra": ["Kayttaja ja paakayttaja vahvistavat kaynnistys ennen huoltoa. Sailytys käsitellään vaatimusten mukaan.", "Tuki-ilmoitus lähetetään sahkoposti kautta. Kayttotunnus ja lokitus tarkistetaan kayttajaohje avulla ennen huoltoa. Dokumentaation omistaja tarkistaa ohjeen kuuden kuukauden välein ja arkkitehtuurimuutoksen jälkeen. Jokaisella vaiheella on oltava ennakkoehto, toiminto, odotettu tulos ja turvallinen keskeytyskohta."]
    },
]


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for index, line in enumerate(text.splitlines() or [""]):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for cell, header in zip(table.rows[0].cells, headers):
        shade(cell, NAVY)
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            if row_index % 2 == 0:
                shade(cell, STRIPE)
            set_cell_text(cell, text, size=8)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_callout(document, label, text):
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, PALE)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    heading = paragraph.add_run(label)
    heading.bold = True
    heading.font.name = "Aptos"
    heading.font.size = Pt(8)
    heading.font.color.rgb = RGBColor.from_string(TEAL)
    body = cell.add_paragraph(text)
    body.paragraph_format.space_after = Pt(2)
    for run in body.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8.3)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_diagram(document, nodes):
    table = document.add_table(rows=1, cols=len(nodes) * 2 - 1)
    table.autofit = True
    for index, (title, details) in enumerate(nodes):
        cell = table.cell(0, index * 2)
        shade(cell, STRIPE)
        set_cell_text(cell, f"{title}\n{details}", bold=True, size=7.5)
        if index < len(nodes) - 1:
            set_cell_text(table.cell(0, index * 2 + 1), "->", bold=True, color=TEAL, size=8)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def configure(document, title):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(8.6)
    normal.font.color.rgb = RGBColor.from_string("263746")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.02
    for style_name, size, color in [("Title", 23, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 11, TEAL)]:
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(1)
        style.paragraph_format.space_after = Pt(4)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run(f"EDGE PLATFORM  /  ENGINEERING                                      {title.upper()}")
    run.bold = True
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("INTERNAL OPERATIONS  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_item(document, item):
    if item[0] == "p":
        document.add_paragraph(item[1])
    elif item[0] == "callout":
        add_callout(document, item[1], item[2])
    elif item[0] == "table":
        add_table(document, item[1], item[2])
    elif item[0] == "bullets":
        for text in item[1]:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.add_run(text)
    elif item[0] == "diagram":
        add_diagram(document, item[1])


def add_section(document, page, extra=()):
    heading = document.add_heading(page["title"], level=1)
    heading.paragraph_format.keep_with_next = True
    subtitle = document.add_paragraph(page["subtitle"])
    subtitle.paragraph_format.space_after = Pt(8)
    for run in subtitle.runs:
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("607382")
    for item in page["items"]:
        add_item(document, item)
    for paragraph in extra:
        document.add_paragraph(paragraph)


def body_text(document):
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def write_english():
    document = Document()
    configure(document, "Operations handbook")
    for index, page in enumerate(PAGES):
        if index:
            document.add_page_break()
        add_section(document, page, SUPPLEMENTAL_PARAGRAPHS[index])
    ENGLISH_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(ENGLISH_OUTPUT)
    verify(ENGLISH_OUTPUT, TYPO_LEDGER)


def write_finnish():
    if len(FINNISH_PAGES) != 10 or len(FINNISH_TYPO_LEDGER) != 40:
        raise ValueError("Expected ten Finnish sections and forty deliberate spelling mistakes.")
    document = Document()
    configure(document, "Käyttökäsikirja")
    for index, page in enumerate(FINNISH_PAGES):
        if index:
            document.add_page_break()
        add_section(document, page, page["extra"])
    FINNISH_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(FINNISH_OUTPUT)
    verify(FINNISH_OUTPUT, FINNISH_TYPO_LEDGER)


def verify(path, ledger):
    document = Document(path)
    tokens = Counter(re.findall(r"\w+", body_text(document).casefold()))
    counts = {word: tokens[word.casefold()] for word in ledger}
    mismatch = {word: count for word, count in counts.items() if count != 1}
    if mismatch:
        raise ValueError(f"{path.name}: typo ledger mismatch: {mismatch}")
    page_breaks = document.element.body.xpath('.//w:br[@w:type="page"]')
    if len(page_breaks) != 9:
        raise ValueError(f"{path.name}: expected 9 explicit section breaks, got {len(page_breaks)}")
    print(f"{path.name}: 10 sections, {sum(counts.values())} deliberate misspellings verified.")


if __name__ == "__main__":
    write_english()
    write_finnish()