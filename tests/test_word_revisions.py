from copy import deepcopy
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document as Word
from lxml import etree as ET
from test_annotations import finding, make_word

from spellcheck_tool.models import DocumentError
from spellcheck_tool.pipeline import export, load_document
from spellcheck_tool.word import NS, package, parse, serialize, w


def resolved_document(data, accept):
    """Simulate Word accepting/rejecting text revisions, then read the DOCX normally."""
    files = package(data)
    root = parse(files["word/document.xml"])
    for revision in root.xpath(".//w:ins | .//w:del", namespaces=NS):
        if (revision.tag == w("ins")) == accept:
            for run in list(revision):
                for node in run.iter(w("delText")):
                    node.tag = w("t")
                revision.addprevious(run)
        revision.getparent().remove(revision)
    files["word/document.xml"] = serialize(root)
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return Word(stream)


def tracked_export(document, findings):
    return export(
        document,
        findings,
        [{"id": issue.id, "correction": issue.suggestions[0]} for issue in findings],
        track_changes=True,
    )


def test_word_revisions_accept_reject_and_preserve_formatting_and_package():
    original = make_word()
    document = load_document("sample.docx", original)
    body = next(s for s in document.segments if "recieve" in s.text)
    table = next(s for s in document.segments if "speling" in s.text)
    issues = [finding(body, "recieve", "receive"), finding(table, "speling", "spelling")]
    output = tracked_export(document, issues)
    before, after = package(original), package(output)
    assert before.keys() == after.keys()
    for name in before:
        if name != "word/document.xml":
            assert before[name] == after[name], name
    root = parse(after["word/document.xml"])
    assert root.xpath(".//w:del//w:delText/text()", namespaces=NS) == ["rec", "ieve", "speling"]
    assert root.xpath(".//w:ins//w:t/text()", namespaces=NS) == ["receive", "spelling"]
    revisions = root.xpath(".//w:del | .//w:ins", namespaces=NS)
    assert len({r.get(w("id")) for r in revisions}) == 4
    assert {r.get(w("author")) for r in revisions} == {"Spellcheck Tool"}
    assert all(r.get(w("date")).endswith("Z") for r in revisions)
    assert not root.xpath(".//w:del//w:t | .//w:ins//w:delText", namespaces=NS)
    accepted = resolved_document(output, accept=True)
    rejected = resolved_document(output, accept=False)
    assert accepted.paragraphs[1].text == "We receive the report. We recieve another report."
    assert accepted.tables[0].cell(0, 0).text == "A spelling mistake in a table."
    assert "".join(r.text for r in accepted.paragraphs[1].runs if r.bold) == "receive"
    original_word = Word(BytesIO(original))
    assert [p.text for p in rejected.paragraphs] == [p.text for p in original_word.paragraphs]
    assert rejected.tables[0].cell(0, 0).text == original_word.tables[0].cell(0, 0).text
    assert "".join(r.text for r in rejected.paragraphs[1].runs if r.bold) == "rec"


def test_word_revisions_multiple_adjacent_edits_and_normalized_offsets():
    word = Word()
    word.add_paragraph("A \ufb01ne mis\u00adspeling:\tspelingrecieve\n25\u00a0kg.")
    stream = BytesIO()
    word.save(stream)
    document = load_document("sample.docx", stream.getvalue())
    segment = document.segments[0]
    issues = [
        finding(segment, "misspeling", "misspelling"),
        finding(segment, "speling", "spelling", 1),
        finding(segment, "recieve", "receive"),
        finding(segment, "25 kg", "twenty-five\tkg\r\n"),
    ]
    output = tracked_export(document, issues)
    assert resolved_document(output, accept=True).paragraphs[0].text == (
        "A \ufb01ne misspelling:\tspellingreceive\ntwenty-five\tkg\n."
    )
    assert resolved_document(output, accept=False).paragraphs[0].text == word.paragraphs[0].text


@pytest.mark.parametrize("existing_revisions", [False, True])
def test_word_revisions_retain_existing_comments_and_revisions(existing_revisions):
    document = load_document("sample.docx", make_word())
    body = next(s for s in document.segments if "recieve" in s.text)
    issue = finding(body, "recieve", "receive")
    first = export(
        document, [issue], [{"id": issue.id, "correction": "receive"}],
        track_changes=existing_revisions,
    )
    document = load_document("sample.docx", first)
    table = next(s for s in document.segments if "speling" in s.text)
    output = tracked_export(document, [finding(table, "speling", "spelling")])
    before, after = package(first), package(output)
    for name in before:
        if name != "word/document.xml":
            assert after[name] == before[name], name
    old_root, new_root = parse(before["word/document.xml"]), parse(after["word/document.xml"])
    assert ET.tostring(old_root.find("w:body", NS)[1]) == ET.tostring(new_root.find("w:body", NS)[1])
    ids = new_root.xpath(".//w:ins/@w:id | .//w:del/@w:id", namespaces=NS)
    assert len(ids) == len(set(ids))


def test_word_tracked_export_validates_choices_and_empty_export():
    original = make_word()
    document = load_document("sample.docx", original)
    assert export(document, [], [], track_changes=True) == original
    segment = next(s for s in document.segments if "recieve" in s.text)
    issue = finding(segment, "recieve", "receive")
    overlap = deepcopy(issue)
    overlap.id = "overlap"
    with pytest.raises(DocumentError, match="overlap"):
        tracked_export(document, [issue, overlap])
    with pytest.raises(DocumentError, match="available"):
        export(
            document, [issue], [{"id": issue.id, "correction": "unapproved"}], track_changes=True
        )
    issue.original = "does not match"
    with pytest.raises(DocumentError, match="no longer matches"):
        tracked_export(document, [issue])
