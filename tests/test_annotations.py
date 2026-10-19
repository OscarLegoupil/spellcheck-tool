from copy import deepcopy
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf as fitz
import pytest
from docx import Document as Word

from spellcheck_tool.models import DocumentError, Finding
from spellcheck_tool.pdf import apply_pdf_changes, pdf_correction_targets
from spellcheck_tool.pipeline import export, load_document
from spellcheck_tool.word import NS, parse, serialize, w


def make_pdf():
    pdf = fitz.open()
    p = pdf.new_page()
    p.insert_text((50, 80), "We recieve the report. We recieve another report.", fontsize=13)
    p.insert_text((50, 130), "Second line stays unchanged.", fontsize=13)
    old = p.add_text_annot((40, 40), "Existing note")
    old.update()
    data = pdf.tobytes()
    pdf.close()
    return data


def make_word():
    doc = Word()
    doc.add_heading("Review fixture", 0)
    p = doc.add_paragraph()
    p.add_run("We ")
    p.add_run("rec").bold = True
    p.add_run("ieve")
    p.add_run(" the report. We recieve another report.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "A speling mistake in a table."
    doc.sections[0].header.paragraphs[0].text = "Header retained"
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def finding(segment, word, correction, occurrence=0):
    start = -1
    for _ in range(occurrence + 1):
        start = segment.text.index(word, start + 1)
    return Finding(
        segment.id, start, start + len(word), word, [correction], "TEST", "spelling", "test"
    )


def test_pdf_correct_occurrence_and_content_preservation():
    original = make_pdf()
    document = load_document("sample.pdf", original)
    segment = next(s for s in document.segments if "recieve" in s.text)
    f = finding(segment, "recieve", "receive", 1)
    result = export(document, [f], [{"id": f.id, "correction": "receive"}])
    with (
        fitz.open(stream=original, filetype="pdf") as before,
        fitz.open(stream=result, filetype="pdf") as after,
    ):
        assert before[0].get_text() == after[0].get_text()
        assert before[0].read_contents() == after[0].read_contents()
        # Transparency can change MuPDF edge antialiasing by 1–2 levels.
        a = before[0].get_pixmap(annots=False).samples
        b = after[0].get_pixmap(annots=False).samples
        assert len(a) == len(b)
        assert max(abs(x - y) for x, y in zip(a, b)) <= 2
        page = after[0]
        notes = list(page.annots())
        assert [n.info["content"] for n in notes] == ["Existing note", "receive"]
        assert notes[-1].rect.x0 > after[0].search_for("recieve")[1].x0 - 10
        assert notes[-1].opacity == pytest.approx(0.28)
        assert notes[-1].info["title"] == ""


def test_pdf_visual_correction_replaces_only_selected_occurrence():
    original = make_pdf()
    document = load_document("sample.pdf", original)
    segment = next(s for s in document.segments if "recieve" in s.text)
    issue = finding(segment, "recieve", "receive", 1)
    target = pdf_correction_targets(document, [issue])[issue.id]
    assert target["supported"]
    corrected = apply_pdf_changes(document, [issue], {issue.id: "receive"})
    with (
        fitz.open(stream=original, filetype="pdf") as before,
        fitz.open(stream=corrected, filetype="pdf") as after,
    ):
        assert len(before[0].search_for("recieve")) == 2
        assert len(after[0].search_for("recieve")) == 1
        assert len(after[0].search_for("receive")) == 1
        assert after[0].search_for("receive")[0].width == pytest.approx(
            before[0].search_for("recieve")[1].width, abs=1
        )
        assert any(annotation.info["content"] == "Existing note" for annotation in after[0].annots())


def test_docx_split_runs_and_multiple_annotations():
    original = make_word()
    document = load_document("sample.docx", original)
    assert document.name == "sample.docx"
    seg = next(s for s in document.segments if "recieve" in s.text)
    f1 = finding(seg, "recieve", "receive")
    f2 = finding(seg, "recieve", "receive", 1)
    ts = next(s for s in document.segments if "speling" in s.text)
    f3 = finding(ts, "speling", "spelling")
    result = export(
        document, [f1, f2, f3], [{"id": f.id, "correction": f.suggestions[0]} for f in [f1, f2, f3]]
    )
    with ZipFile(BytesIO(original)) as a, ZipFile(BytesIO(result)) as b:
        ra, rb = parse(a.read("word/document.xml")), parse(b.read("word/document.xml"))
        assert "".join(ra.xpath(".//w:t/text()", namespaces=NS)) == "".join(
            rb.xpath(".//w:t/text()", namespaces=NS)
        )
        assert b.read("word/header1.xml") == a.read("word/header1.xml")
        assert b.read("word/styles.xml") == a.read("word/styles.xml")
        assert len(rb.xpath(".//w:commentRangeStart", namespaces=NS)) == 3
        assert len(rb.xpath(".//w:commentRangeEnd", namespaces=NS)) == 3
        assert len(rb.xpath(".//w:commentReference", namespaces=NS)) == 3
        assert parse(b.read("word/comments.xml")).xpath(".//w:t/text()", namespaces=NS) == [
            "receive",
            "receive",
            "spelling",
        ]
        bold_text = rb.xpath(".//w:r[w:rPr/w:b]/w:t/text()", namespaces=NS)
        assert "".join(bold_text) == "rec"
        assert len(rb.xpath('.//w:shd[@w:fill="FFF2CC"]', namespaces=NS)) >= 3
    assert any("header" in x for x in document.warnings)
    Word(BytesIO(result))


def test_existing_comments_retained_and_paragraph_not_reannotated():
    document = load_document("sample.docx", make_word())
    seg = next(s for s in document.segments if "recieve" in s.text)
    f = finding(seg, "recieve", "receive")
    first = export(document, [f], [{"id": f.id, "correction": "receive"}])
    again = load_document("sample.docx", first)
    assert not any("recieve" in s.text for s in again.segments)
    table = next(s for s in again.segments if "speling" in s.text)
    f2 = finding(table, "speling", "spelling")
    second = export(again, [f2], [{"id": f2.id, "correction": "spelling"}])
    with ZipFile(BytesIO(second)) as z:
        comments = parse(z.read("word/comments.xml"))
        assert comments.xpath(".//w:t/text()", namespaces=NS) == ["receive", "spelling"]
        assert len({c.get(w("id")) for c in comments}) == 2


def test_overlapping_findings_are_rejected():
    document = load_document("sample.pdf", make_pdf())
    seg = document.segments[0]
    f = finding(seg, "recieve", "receive")
    other = deepcopy(f)
    other.id = "other"
    with pytest.raises(DocumentError, match="overlap"):
        export(document, [f, other], [{"id": x.id, "correction": "receive"} for x in [f, other]])


def test_arbitrary_replacement_rejected():
    document = load_document("sample.pdf", make_pdf())
    f = finding(document.segments[0], "recieve", "receive")
    with pytest.raises(DocumentError, match="available"):
        export(document, [f], [{"id": f.id, "correction": "a different sentence"}])


def test_no_selection_docx_is_byte_identical():
    original = make_word()
    document = load_document("sample.docx", original)
    assert export(document, [], []) == original


def test_image_only_pdf_is_reported():
    pdf = fitz.open()
    pdf.new_page()
    data = pdf.tobytes()
    pdf.close()
    with pytest.raises(DocumentError, match="selectable text"):
        load_document("scan.pdf", data)


def test_encrypted_pdf_rejected():
    with fitz.open(stream=make_pdf(), filetype="pdf") as pdf:
        data = pdf.tobytes(
            encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="password"
        )
    with pytest.raises(DocumentError, match="Password"):
        load_document("locked.pdf", data)


def test_ligature_and_soft_hyphen_offsets_in_word():
    doc = Word()
    doc.add_paragraph("A ﬁne mis\u00adspeling in text.")
    stream = BytesIO()
    doc.save(stream)
    document = load_document("sample.docx", stream.getvalue())
    seg = document.segments[0]
    f = finding(seg, "misspeling", "misspelling")
    output = export(document, [f], [{"id": f.id, "correction": "misspelling"}])
    with ZipFile(BytesIO(output)) as z:
        text = "".join(parse(z.read("word/document.xml")).xpath(".//w:t/text()", namespaces=NS))
    assert text == "A ﬁne mis\u00adspeling in text."


def test_revision_and_hyperlink_coverage():
    original = make_word()
    with ZipFile(BytesIO(original)) as z:
        files = {n: z.read(n) for n in z.namelist()}
    root = parse(files["word/document.xml"])
    p = root.xpath(".//w:p", namespaces=NS)[1]
    p[1].tag = w("hyperlink")
    files["word/document.xml"] = serialize(root)
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as z:
        for n, b in files.items():
            z.writestr(n, b)
    document = load_document("sample.docx", output.getvalue())
    assert any("links" in w for w in document.warnings)
    assert not any("recieve" in s.text for s in document.segments)
